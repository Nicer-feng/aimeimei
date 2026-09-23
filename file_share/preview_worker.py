"""Executed only inside the network/filesystem sandbox; no application imports."""
import csv
import json
import resource
import subprocess
import sys
import zipfile
from pathlib import Path

resource.setrlimit(resource.RLIMIT_AS, (768*1024**2,768*1024**2))
resource.setrlimit(resource.RLIMIT_CPU, (25,30))
resource.setrlimit(resource.RLIMIT_FSIZE, (25*1024**2,25*1024**2))
resource.setrlimit(resource.RLIMIT_NOFILE, (128,128))
resource.setrlimit(resource.RLIMIT_NPROC, (64,64))
source = Path(sys.argv[1])
ext = source.suffix.lower()
if ext in ('.docx','.xlsx'):
    with zipfile.ZipFile(source) as archive:
        info=archive.infolist()
        if len(info)>5000 or sum(x.file_size for x in info)>80*1024**2:
            raise ValueError('Expanded document too large')
        if any(x.flag_bits & 1 for x in info):raise ValueError('Encrypted file')

if ext in ('.doc','.docx'):
    profile=Path('/work/profile/user');profile.mkdir(parents=True)
    (profile/'registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item><item oor:path="/org.openoffice.Office.Common/Load"><prop oor:name="UpdateDocMode" oor:op="fuse"><value>0</value></prop></item></oor:items>''')
    subprocess.run(['/usr/bin/libreoffice','-env:UserInstallation=file:///work/profile','--headless','--nologo','--nodefault','--norestore','--convert-to','pdf:writer_pdf_Export','--outdir','/work',str(source)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=30)
    result=source.with_suffix('.pdf')
    if not result.exists() or result.stat().st_size>20*1024**2:raise ValueError('Conversion failed')
    result.rename('/work/result.pdf')
else:
    sheets=[];budget=500000;cells=0;truncated=False
    def rows(values):
        global budget,cells,truncated
        result=[]
        for n,row in enumerate(values):
            if n>=200 or cells>=30000 or budget<=0:truncated=True;break
            out=[]
            for value in row[:50]:
                text='' if value is None else str(value)
                if len(text)>500:truncated=True
                text=text[:min(500,max(0,budget))];budget-=len(text);cells+=1;out.append(text)
            if len(row)>50:truncated=True
            result.append(out)
        return result
    if ext=='.xlsx':
        import openpyxl
        book=openpyxl.load_workbook(source,read_only=True,data_only=True,keep_links=False)
        try:
            for sheet in book.worksheets[:20]:
                if sheet.sheet_state!='visible':continue
                if (sheet.max_row or 0)>200 or (sheet.max_column or 0)>50:truncated=True
                sheets.append({'name':sheet.title,'rows':rows(sheet.iter_rows(max_row=min(sheet.max_row or 200,201),max_col=min(sheet.max_column or 50,50),values_only=True))})
            truncated=truncated or len(book.worksheets)>20
        finally:book.close()
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(str(source),on_demand=True)
        try:
            for n in range(min(20,book.nsheets)):
                sheet=book.sheet_by_index(n)
                if sheet.visibility:continue
                if sheet.nrows>200 or sheet.ncols>50:truncated=True
                def values():
                    for r in range(min(sheet.nrows,201)):
                        result=[]
                        for cell in sheet.row(r)[:50]:
                            value=cell.value
                            if cell.ctype==xlrd.XL_CELL_DATE:
                                try:value=xlrd.xldate_as_datetime(value,book.datemode).isoformat(sep=' ')
                                except ValueError:pass
                            result.append(value)
                        yield result
                sheets.append({'name':sheet.name,'rows':rows(values())});book.unload_sheet(n)
            truncated=truncated or book.nsheets>20
        finally:book.release_resources()
    elif ext=='.csv':
        with source.open(encoding='utf-8-sig',errors='replace',newline='') as stream:
            sheets.append({'name':'CSV','rows':rows(csv.reader(stream))})
    else:raise ValueError('Unsupported extension')
    Path('/work/result.json').write_text(json.dumps({'kind':'sheets','sheets':sheets,'truncated':truncated},ensure_ascii=False))
